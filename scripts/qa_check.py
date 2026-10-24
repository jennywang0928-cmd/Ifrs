"""交檔前自我檢查：python3 qa_check.py <excel>。列出所有可疑之處，沒有問題時印出 OK。"""
import os, re, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from email_fix import bad_tld  # noqa: E402
from collections import Counter
from openpyxl import load_workbook

EXPECTED = ['公司代號', '公司簡稱', '公司名稱', '市場', '產業', '建議寄送 email（對應聯絡人）', 'email 來源',
            '聯絡人部門', '聯絡人職稱', '聯絡人姓名', '稱呼姓', '稱呼職稱',
            '永續報告書聯絡 email（申報資料）', '報告書內文聯絡 email', '公司信箱 (發言人/IR)', '永續報告書聯絡資訊原文',
            '聯絡資訊在報告書第幾頁', '報告期間', '永續報告書網址', '公司網址', '聯絡人電話（報告書）', '公司總機電話', '需確認事項']
EMAIL = re.compile(r'^[a-z0-9._%+\-]+@[a-z0-9\-]+(\.[a-z0-9\-]+)*\.[a-z]{2,}$')
TITLE_WORDS = re.compile(r'委員會|小組|辦公室|部|處|室|窗口|中心|單位|公司')

wb = load_workbook(sys.argv[1])
ws = wb['寄送名單']
rows = list(ws.iter_rows(values_only=True))
h, data = list(rows[0]), rows[1:]
col = {name: i for i, name in enumerate(h)}
issues = []
add = lambda kind, code, msg: issues.append((kind, code, msg))

if h != EXPECTED:
    add('欄位', '-', f'欄位順序與預期不同：{h}')
codes = Counter(r[0] for r in data)
for c, n in codes.items():
    if n > 1:
        add('重複', c, f'公司代號出現 {n} 次')

for i, r in enumerate(data, start=2):
    g = lambda name: r[col[name]] or ''
    code = r[0]
    for v in r:
        if isinstance(v, str) and (v[:1] in '=@' or re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', v)):
            add('格式', code, f'儲存格內容可疑：{v[:30]}')
    f = g('建議寄送 email（對應聯絡人）')
    for e in filter(None, f.split('; ')):
        if not EMAIL.match(e):
            add('email', code, f'F 欄 email 格式錯誤：{e}')
        if bad_tld(e):
            add('email', code, f'F 欄 email 結尾可疑（可能黏到後面的字）：{e}')
        if e[:1] in '._-':
            add('email', code, f'F 欄 email 帳號開頭不完整（可能被截斷）：{e}')
    if not f and '查無 email' not in g('需確認事項'):
        add('email', code, 'F 欄空白但 W 欄沒有標示')
    pool = ' '.join([g('永續報告書聯絡 email（申報資料）'), g('報告書內文聯絡 email'), g('公司信箱 (發言人/IR)')])
    if f and f.split('; ')[0] not in pool:
        add('email', code, f'F 欄 email 不在 M/N/O 欄：{f}')
    name, title, sur, st = g('聯絡人姓名'), g('聯絡人職稱'), g('稱呼姓'), g('稱呼職稱')
    # 「簡處長」「閔副總經理」是報告書只寫姓加職稱，屬正常
    if name and TITLE_WORDS.search(name) and not re.search(r'先生|小姐|女士', name) and not (sur and st and name == sur + st):
        add('姓名', code, f'J 欄看起來不是人名：{name}')
    if name and re.fullmatch(r'[一-鿿]{5,}', name) and name != sur + st:
        add('姓名', code, f'J 欄姓名過長：{name}')
    if '、' in name or '；' in title:
        add('聯絡人', code, f'仍有多位聯絡人／多職稱：{name}｜{title}')
    if st and not sur:
        add('稱呼', code, f'L 欄有值但 K 欄空白：{st}')
    if sur and not name:
        add('稱呼', code, f'K 欄有值但 J 欄空白：{sur}')
    if sur and name and not re.match(r'[A-Za-z]', name) and not name.startswith(sur):
        add('稱呼', code, f'K 欄姓氏和 J 欄姓名對不上：{sur}｜{name}')
    if st and re.search(r'[A-Za-z]', st):
        add('稱呼', code, f'L 欄不應是英文職稱：{st}')
    if sur and re.match(r'[A-Za-z]', name) and st:
        add('稱呼', code, f'英文名不應有稱呼職稱：{name}｜{st}')
    phone = g('聯絡人電話（報告書）')
    if phone and len(re.sub(r'\D', '', phone)) < 8:
        add('電話', code, f'U 欄電話位數過少：{phone}')
    if phone and re.search(r'0800|傳真|fax', phone, re.I):
        add('電話', code, f'U 欄可能不是聯絡電話：{phone}')
    for name_ in ('永續報告書網址', '公司網址'):
        c = ws.cell(i, col[name_] + 1)
        if c.value and not c.hyperlink and '網址無效' not in g('需確認事項') and name_ == '永續報告書網址':
            add('網址', code, f'{name_}沒有超連結且未標示：{c.value}')
        if c.hyperlink and ' ' in c.hyperlink.target:
            add('網址', code, f'超連結含空白：{c.hyperlink.target}')
    lcell = ws.cell(i, col['稱呼職稱'] + 1)
    note = lcell.comment.text if lcell.comment else ''
    if '沒把握' in note and (lcell.fill.fgColor.rgb in (None, '00000000') or '依名字推測' not in g('需確認事項')):
        add('稱呼', code, '依名字推測沒把握，但 L 欄沒有黃底或 W 欄沒有標示')
    if name and not sur and not st and not re.match(r'[A-Za-z]', name) and not TITLE_WORDS.search(name):
        add('稱呼', code, f'有中文姓名但沒有稱呼：{name}')
    w = g('需確認事項')
    if w and ws.cell(i, col['需確認事項'] + 1).fill.fgColor.rgb in (None, '00000000'):
        add('標示', code, f'W 欄有內容但沒有上色：{w}')

kinds = Counter(k for k, _, _ in issues)
print('列數', len(data), '| 問題', len(issues), dict(kinds))
for k, c, m in issues[:200]:
    print(f'[{k}] {c} {m}')
print('L 欄稱呼分布', Counter(r[col['稱呼職稱']] for r in data if r[col['稱呼職稱']]).most_common(40))
if not issues:
    print('OK')
